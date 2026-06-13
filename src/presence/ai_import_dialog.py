"""
ai_import_dialog.py — AI-powered Import & Convert dialog.

Accepts a document (PDF, DOCX, TXT, MD), extracts its text, calls the
Claude API to generate a Markdown slide deck, and writes it into the editor.
Image generation is handled separately via the Insert Image → Generate with AI
dialog (ai_image_dialog.py).
"""
from __future__ import annotations

import asyncio
import json
import logging
import re
import threading
from pathlib import Path

import gi
gi.require_version("Gtk", "4.0")
gi.require_version("Adw", "1")
from gi.repository import Gtk, Adw, GLib

from .app_utils import make_file_filter, make_filter_store
from .session import load_api_keys, save_ai_prefs, load_ai_prefs
from .slides.image_gen import (
    generate_all_images, slide_headline, slide_image_scene,
    DEFAULT_GEMINI_MODEL,
)

log = logging.getLogger(__name__)


_TONE_LABELS = ["Formal", "Engaging", "Captivating", "Energetic"]
_TONE_DESCRIPTIONS = {
    "Formal":      "Formal — clear, structured, professional",
    "Engaging":    "Engaging — conversational, relatable, uses rhetorical questions",
    "Captivating": "Captivating — storytelling-driven, vivid, uses metaphor",
    "Energetic":   "Energetic — punchy, motivating, short sentences, active voice",
}
_LANG_LABELS = ["English", "Dutch", "Turkish"]

_NOTES_LABELS = ["None", "Brief", "Standard"]
_NOTES_SENTENCE_RANGES = {
    "None":     None,
    "Brief":    "1–2",
    "Standard": "3–5",
}

_SYSTEM_PROMPT = (
    "You are an expert presentation designer and strategic communicator. "
    "You transform dense source documents into compelling, well-structured "
    "slide decks written in Markdown. You think in narratives: every deck "
    "you produce opens by establishing context or a problem, builds toward "
    "a solution or insight, and closes with a concrete call to action. "
    "You are concise, precise, and never pad slides with filler content."
)

_ANALYSIS_SYSTEM_PROMPT = (
    "You are an expert presentation strategist. "
    "Analyse the structure and content of documents and recommend the ideal "
    "number of slides for a compelling narrative presentation. "
    "Respond with valid JSON only — no text outside the JSON object."
)


class AIImportDialog(Adw.Dialog):
    def __init__(self, parent_window) -> None:
        super().__init__()
        self.set_title("Import & Convert")
        self.set_content_width(480)
        self._parent_window = parent_window
        self._selected_file: Path | None = None
        self._selected_gfile = None  # GFile kept for portal files without a local path
        self._active_file_dialog = None
        self._converting = False
        self._analysis_timeout_id: int | None = None
        self._document_text: str | None = None
        self._pending_tone: str = ""
        self._pending_language: str = ""
        self._pending_notes: str = "Standard"
        self._pending_claude_key: str = ""
        self._build_ui()

        last = getattr(parent_window, "_last_ai_import_file", None)
        if last and last.is_file():
            self._selected_file = last
            self._file_row.set_subtitle(last.name)

    # ── UI construction ───────────────────────────────────────────────────────

    def _build_ui(self) -> None:
        toolbar_view = Adw.ToolbarView()
        self.set_child(toolbar_view)

        header = Adw.HeaderBar()
        header.set_show_end_title_buttons(False)

        self._cancel_btn = Gtk.Button(label="Cancel")
        self._cancel_btn.connect("clicked", lambda *_: self.close())
        header.pack_start(self._cancel_btn)

        self._back_btn = Gtk.Button(label="Back")
        self._back_btn.connect("clicked", self._on_back_clicked)
        self._back_btn.set_visible(False)
        header.pack_start(self._back_btn)

        toolbar_view.add_top_bar(header)

        self._toast_overlay = Adw.ToastOverlay()
        toolbar_view.set_content(self._toast_overlay)

        content = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=0)
        content.set_margin_top(12)
        content.set_margin_bottom(24)
        content.set_margin_start(12)
        content.set_margin_end(12)
        self._toast_overlay.set_child(content)

        # ── Form area ─────────────────────────────────────────────────────────
        self._form_area = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=0)
        content.append(self._form_area)

        group = Adw.PreferencesGroup()
        self._form_area.append(group)

        self._file_row = Adw.ActionRow(title="Document")
        self._file_row.set_subtitle("No file selected")
        choose_btn = Gtk.Button(label="Choose…")
        choose_btn.add_css_class("flat")
        choose_btn.set_valign(Gtk.Align.CENTER)
        choose_btn.connect("clicked", self._on_choose_file)
        self._file_row.add_suffix(choose_btn)
        group.add(self._file_row)

        self._tone_row = Adw.ComboRow(title="Language tone")
        self._tone_row.set_model(Gtk.StringList.new(_TONE_LABELS))
        group.add(self._tone_row)

        self._lang_row = Adw.ComboRow(title="Output language")
        self._lang_row.set_model(Gtk.StringList.new(_LANG_LABELS))
        group.add(self._lang_row)

        self._notes_row = Adw.ComboRow(title="Speaker notes")
        self._notes_row.set_model(Gtk.StringList.new(_NOTES_LABELS))
        self._notes_row.set_selected(2)  # default: Standard
        group.add(self._notes_row)

        disclosure = Gtk.Label(
            label="Document content will be sent to Anthropic's Claude API for processing."
        )
        disclosure.add_css_class("caption")
        disclosure.add_css_class("dim-label")
        disclosure.set_wrap(True)
        disclosure.set_xalign(0)
        disclosure.set_margin_top(8)
        disclosure.set_margin_start(4)
        self._form_area.append(disclosure)

        # ── Review area (shown after analysis) ───────────────────────────────
        self._review_box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=0)
        self._review_box.set_visible(False)
        content.append(self._review_box)

        review_group = Adw.PreferencesGroup()
        self._review_box.append(review_group)

        self._count_row = Adw.ActionRow(title="Number of slides")
        self._count_row.set_subtitle("Adjust to your preference")
        self._slide_count_spin = Gtk.SpinButton.new_with_range(3, 50, 1)
        self._slide_count_spin.set_value(10)
        self._slide_count_spin.set_numeric(True)
        self._slide_count_spin.set_valign(Gtk.Align.CENTER)
        self._count_row.add_suffix(self._slide_count_spin)
        self._count_row.set_activatable_widget(self._slide_count_spin)
        review_group.add(self._count_row)

        self._reasoning_row = Adw.ActionRow(title="Why this count")
        self._reasoning_row.set_subtitle("…")
        self._reasoning_row.set_subtitle_lines(0)
        review_group.add(self._reasoning_row)

        # ── Progress area ─────────────────────────────────────────────────────
        self._progress_box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=6)
        self._progress_box.set_margin_top(12)
        self._progress_box.set_visible(False)

        self._spinner = Gtk.Spinner()
        self._spinner.set_halign(Gtk.Align.CENTER)
        self._progress_box.append(self._spinner)

        self._progress_label = Gtk.Label(label="Analysing document…")
        self._progress_label.add_css_class("caption")
        self._progress_label.add_css_class("dim-label")
        self._progress_label.set_halign(Gtk.Align.CENTER)
        self._progress_box.append(self._progress_label)

        content.append(self._progress_box)

        # ── Buttons ───────────────────────────────────────────────────────────
        self._convert_btn = Gtk.Button(label="Analyse & Continue")
        self._convert_btn.add_css_class("suggested-action")
        self._convert_btn.set_halign(Gtk.Align.FILL)
        self._convert_btn.set_margin_top(12)
        self._convert_btn.connect("clicked", self._on_convert_clicked)
        content.append(self._convert_btn)

        self._generate_btn = Gtk.Button(label="Generate Slides")
        self._generate_btn.add_css_class("suggested-action")
        self._generate_btn.set_halign(Gtk.Align.FILL)
        self._generate_btn.set_margin_top(12)
        self._generate_btn.connect("clicked", self._on_generate_clicked)
        self._generate_btn.set_visible(False)
        content.append(self._generate_btn)

    # ── Signal handlers ───────────────────────────────────────────────────────

    def _on_choose_file(self, *_) -> None:
        dialog = Gtk.FileDialog()
        dialog.set_title("Choose Document")
        all_files = Gtk.FileFilter()
        all_files.set_name("All files")
        all_files.add_pattern("*")
        dialog.set_filters(make_filter_store(
            make_file_filter(
                "Supported documents", "*.pdf", "*.docx", "*.txt", "*.md"
            ),
            all_files,
        ))
        self._active_file_dialog = dialog
        dialog.open(self._parent_window, None, self._on_file_chosen)

    def _on_file_chosen(self, dialog: Gtk.FileDialog, result) -> None:
        self._active_file_dialog = None
        try:
            gfile = dialog.open_finish(result)
        except GLib.Error:
            return
        self._selected_gfile = gfile
        path_str = gfile.get_path()
        if path_str:
            path = Path(path_str)
            if path.is_file():
                self._selected_file = path
                self._file_row.set_subtitle(path.name)
        else:
            # Portal file without a FUSE path — read via GLib when analysing.
            self._selected_file = None
            self._file_row.set_subtitle(gfile.get_basename() or "selected file")

    def _on_back_clicked(self, *_) -> None:
        self._enter_form_state()

    def _on_convert_clicked(self, *_) -> None:
        claude_key, _gemini_key = load_api_keys()

        if self._selected_file is None and self._selected_gfile is None:
            self._show_toast("Please select a file first.")
            return
        if not claude_key:
            self._show_toast("Claude API key not set — open Settings to add it.")
            return

        self._pending_tone = _TONE_DESCRIPTIONS[
            _TONE_LABELS[self._tone_row.get_selected()]
        ]
        self._pending_language = _LANG_LABELS[self._lang_row.get_selected()]
        self._pending_notes = _NOTES_LABELS[self._notes_row.get_selected()]
        self._pending_claude_key = claude_key

        self._enter_analyzing_state()

        threading.Thread(
            target=self._run_analysis,
            args=(self._selected_file, claude_key),
            daemon=True,
        ).start()

    def _on_generate_clicked(self, *_) -> None:
        if self._document_text is None:
            self._show_toast("Document text not available — please start over.")
            return
        slide_count = int(self._slide_count_spin.get_value())
        self._enter_generating_state()
        threading.Thread(
            target=self._run_conversion,
            args=(
                self._document_text,
                self._pending_claude_key,
                self._pending_tone,
                self._pending_language,
                slide_count,
                self._pending_notes,
            ),
            daemon=True,
        ).start()

    # ── State management ──────────────────────────────────────────────────────

    def _enter_form_state(self) -> None:
        self._cancel_btn.set_visible(True)
        self._back_btn.set_visible(False)
        self._form_area.set_visible(True)
        self._form_area.set_sensitive(True)
        self._review_box.set_visible(False)
        self._progress_box.set_visible(False)
        self._convert_btn.set_visible(True)
        self._convert_btn.set_sensitive(True)
        self._generate_btn.set_visible(False)
        self._spinner.stop()
        self._converting = False

    def _enter_analyzing_state(self) -> None:
        self._form_area.set_sensitive(False)
        self._progress_label.set_label("Analysing document…")
        self._progress_box.set_visible(True)
        self._convert_btn.set_sensitive(False)
        self._spinner.start()
        self._converting = True
        self._analysis_timeout_id = GLib.timeout_add(
            90_000, self._on_analysis_timeout
        )

    def _on_analysis_timeout(self) -> bool:
        self._analysis_timeout_id = None
        if self._converting:
            self._on_error(
                "Analysis timed out — check your network connection and API key"
            )
        return GLib.SOURCE_REMOVE

    def _enter_review_state(self) -> None:
        self._cancel_btn.set_visible(False)
        self._back_btn.set_visible(True)
        self._form_area.set_visible(False)
        self._form_area.set_sensitive(True)
        self._review_box.set_visible(True)
        self._progress_box.set_visible(False)
        self._convert_btn.set_visible(False)
        self._generate_btn.set_visible(True)
        self._generate_btn.set_sensitive(True)
        self._slide_count_spin.set_sensitive(True)
        self._spinner.stop()
        self._converting = False

    def _enter_generating_state(self) -> None:
        self._generate_btn.set_sensitive(False)
        self._slide_count_spin.set_sensitive(False)
        self._progress_label.set_label("Generating slides…")
        self._progress_box.set_visible(True)
        self._spinner.start()
        self._converting = True

    # ── Background work ───────────────────────────────────────────────────────

    def _run_analysis(self, file_path: Path | None, claude_key: str) -> None:
        try:
            if file_path is not None:
                document_text = _extract_text(file_path)
            else:
                gfile = self._selected_gfile
                if gfile is None:
                    GLib.idle_add(self._on_error, "No file selected.")
                    return
                document_text = _extract_gfile_text(gfile)
        except Exception as e:
            label = file_path.name if file_path else (
                self._selected_gfile.get_basename() if self._selected_gfile else "file"
            )
            GLib.idle_add(self._on_error, f"Could not read {label}: {e}")
            return

        self._document_text = document_text

        try:
            import anthropic
        except Exception as _exc:
            GLib.idle_add(
                self._on_error,
                f"anthropic package not installed or failed to import: {_exc}",
            )
            return

        try:
            user_prompt = (
                "Analyse this document and recommend the ideal number of slides "
                "for a presentation.\n\n"
                "Consider:\n"
                "- Number of distinct topics or sections\n"
                "- Depth and density of content per section\n"
                "- Need for opening, context, key insights, evidence, and a call to action\n"
                "- Avoid padding (too many slides) and avoid rushing (too few)\n\n"
                "Respond with ONLY this JSON (no markdown fences, no extra text):\n"
                "{\n"
                '  "slide_count": <integer, 5–30>,\n'
                '  "reasoning": "<2–3 sentence explanation of why this number best serves the narrative>"\n'
                "}\n\n"
                f"DOCUMENT:\n{document_text}"
            )
            client = anthropic.Anthropic(api_key=claude_key, timeout=60.0)
            message = client.messages.create(
                model="claude-haiku-4-5-20251001",
                max_tokens=300,
                system=_ANALYSIS_SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_prompt}],
            )
            raw = message.content[0].text.strip()
            if raw.startswith("```"):
                raw = raw.split("```")[1]
                if raw.startswith("json"):
                    raw = raw[4:]
            data = json.loads(raw)
            slide_count = max(3, min(50, int(data["slide_count"])))
            reasoning = str(data.get("reasoning", ""))
        except Exception as e:
            log.warning("Analysis call failed, using default slide count: %s", e)
            slide_count = 10
            reasoning = (
                "Could not analyse the document automatically — using a default of 10 slides. "
                "Adjust the count to fit your needs."
            )

        GLib.idle_add(self._on_analysis_done, slide_count, reasoning)

    def _run_conversion(
        self,
        document_text: str,
        claude_key: str,
        tone: str,
        language: str,
        slide_count: int,
        notes_level: str = "Standard",
    ) -> None:
        GLib.idle_add(self._set_progress_label, "Generating slides…")

        try:
            import anthropic
        except Exception as _exc:
            GLib.idle_add(
                self._on_error,
                f"anthropic package not installed or failed to import: {_exc}",
            )
            return

        sentence_range = _NOTES_SENTENCE_RANGES[notes_level]
        if sentence_range is None:
            notes_rule = (
                "- Do NOT include ^^^ or any speaker notes — headlines only."
            )
            notes_example = "# This Is the Slide Headline"
        else:
            notes_rule = (
                f"- Immediately after the headline, write ^^^ on its own line\n"
                f"- After ^^^, write the speaker notes: {sentence_range} sentences of flowing "
                f"prose, no bullet points, no dashes, no markdown formatting — plain continuous "
                f"text written to be read aloud on a teleprompter\n"
                f"- Do NOT add any bullet body between the headline and ^^^"
            )
            notes_example = (
                "# This Is the Slide Headline\n\n"
                "^^^\n\n"
                "First sentence for the teleprompter. Second sentence continues the thought."
                + (" Third sentence closes the point naturally." if sentence_range != "1–2" else "")
            )

        try:
            user_prompt = f"""Convert the document below into a Markdown slide deck of exactly {slide_count} slides.

FORMATTING RULES — follow exactly:
- Separate slides with: ---
- Slide headline: # Headline text (one per slide, max 8 words)
{notes_rule}

EXAMPLE of one correctly formatted slide:

{notes_example}

NARRATIVE STRUCTURE — mandatory, exactly {slide_count} slides total:
1. Slide 1: Opening — title of the deck
2. 2–3 context slides: establish why the topic matters
3. The bulk of slides: key insights or recommendations, one idea per slide
4. 1–2 evidence slides if appropriate: data, examples, or quotes from the source
5. Slide {slide_count}: a single, specific call to action

Distribute the {slide_count} slides proportionally across this structure.
Generate EXACTLY {slide_count} slides separated by --- — no more, no less.

TONE: {tone}
OUTPUT LANGUAGE: {language} — write every word of the output in {language}, including headlines and speaker notes.

If the document does not contain enough content for a full narrative, synthesise logically from what is present rather than inventing facts.

---
DOCUMENT:
{document_text}"""

            client = anthropic.Anthropic(api_key=claude_key, timeout=300.0)
            chunks: list[str] = []
            slides_seen = 0
            chars_since_update = 0

            with client.messages.stream(
                model="claude-sonnet-4-6",
                max_tokens=8000,
                system=_SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_prompt}],
            ) as stream:
                for text in stream.text_stream:
                    chunks.append(text)
                    chars_since_update += len(text)
                    slides_seen += text.count("\n---\n")
                    if chars_since_update >= 150:
                        chars_since_update = 0
                        label = (
                            f"Generating slides… {slides_seen} of {slide_count}"
                            if slides_seen > 0
                            else "Generating slides…"
                        )
                        GLib.idle_add(self._set_progress_label, label)

            markdown_text = "".join(chunks)

        except Exception as e:
            GLib.idle_add(self._on_error, f"Claude API error: {e}")
            return

        GLib.idle_add(self._on_done, markdown_text)

    # ── GLib main-thread callbacks ────────────────────────────────────────────

    def _on_analysis_done(self, slide_count: int, reasoning: str) -> bool:
        if self._analysis_timeout_id is not None:
            GLib.source_remove(self._analysis_timeout_id)
            self._analysis_timeout_id = None
        self._slide_count_spin.set_value(slide_count)
        self._reasoning_row.set_subtitle(GLib.markup_escape_text(reasoning))
        self._enter_review_state()
        return GLib.SOURCE_REMOVE

    def _on_done(self, markdown_text: str) -> bool:
        self._converting = False
        win = self._parent_window
        if self._selected_file is not None:
            win._last_ai_import_file = self._selected_file
        win._editor.set_text(markdown_text)
        win._modified = True
        base = win._file_path.name if win._file_path else "Untitled"
        win._set_title(base + " •")
        win._sidebar.update_from_text(markdown_text)
        self.close()
        return GLib.SOURCE_REMOVE

    def _on_error(self, message: str) -> bool:
        if self._document_text is not None:
            self._enter_review_state()
        else:
            self._enter_form_state()
        dlg = Adw.AlertDialog(heading="Error", body=message)
        dlg.add_response("ok", "OK")
        dlg.set_default_response("ok")
        dlg.set_close_response("ok")
        dlg.present(self._parent_window)
        return GLib.SOURCE_REMOVE

    def _set_progress_label(self, label: str) -> bool:
        self._progress_label.set_label(label)
        return GLib.SOURCE_REMOVE

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _show_toast(self, message: str, timeout: int = 5) -> None:
        toast = Adw.Toast(title=message)
        toast.set_timeout(timeout)
        self._toast_overlay.add_toast(toast)


# ── Module-level helpers ──────────────────────────────────────────────────────

def generate_missing_images(parent_window) -> None:
    """
    Public entry point called from the window action.

    Scans the current editor Markdown for gemini image placeholders whose
    files don't exist yet and generates only the missing ones.  If the deck
    has no image placeholders at all, placeholders are inserted for every
    slide that has a headline and then all images are generated.
    """
    _, gemini_key = load_api_keys()
    if not gemini_key:
        parent_window._show_toast(
            "Gemini API key not set — open Settings to add it."
        )
        return

    if not parent_window._file_path:
        parent_window._show_toast(
            "Save the file first — images are written to assets/ next to the .md file."
        )
        return

    markdown = parent_window._editor.get_text()
    assets_dir = parent_window._file_path.parent / "assets"

    pattern = re.compile(
        r'!\[([^|\]]+)\|[^\]]*\]\(assets/gemini_(\d+)\.jpg\)'
    )

    has_any_placeholder = bool(pattern.search(markdown))

    if not has_any_placeholder:
        markdown, missing = _insert_image_placeholders(markdown)
        if not missing:
            parent_window._show_toast("No slides with headings found.")
            return
        parent_window._editor.set_text(markdown)
        parent_window._modified = True
    else:
        missing = []
        for m in pattern.finditer(markdown):
            scene = m.group(1).strip()
            idx = int(m.group(2))
            if not (assets_dir / f"gemini_{idx}.jpg").exists():
                missing.append((scene, idx))

        if not missing:
            parent_window._show_toast("All slide images are already present.")
            return

    prefs = load_ai_prefs()
    image_style = prefs.get("image_style", "Photorealistic")
    gemini_model = prefs.get("gemini_model", DEFAULT_GEMINI_MODEL)

    parent_window._show_toast(
        f"Generating {len(missing)} missing image{'s' if len(missing) != 1 else ''}…", 0
    )
    log.debug("Missing images: %s (style: %s)", missing, image_style)

    threading.Thread(
        target=_run_image_generation,
        args=(None, gemini_key, str(assets_dir),
              parent_window, missing, image_style, gemini_model),
        daemon=True,
    ).start()


def _insert_image_placeholders(markdown: str) -> tuple[str, list[tuple[str, int]]]:
    """Insert gemini image placeholders into a deck that has none.

    Splits on ``---`` slide separators, finds the first ``# heading`` in each
    slide, and inserts ``![headline|left/right|50|nogradient](assets/gemini_N.jpg)``
    immediately after it.  Returns (updated_markdown, [(scene, index), ...]).
    """
    slides = re.split(r'\n[ \t]*---[ \t]*\n', markdown)
    tasks: list[tuple[str, int]] = []
    new_slides: list[str] = []

    for i, slide in enumerate(slides):
        headline = slide_headline(slide)
        if headline:
            alignment = "left" if i % 2 == 0 else "right"
            placeholder = (
                f"![{headline}|{alignment}|50|nogradient]"
                f"(assets/gemini_{i}.jpg)"
            )
            lines = slide.splitlines()
            new_lines: list[str] = []
            inserted = False
            for line in lines:
                new_lines.append(line)
                if not inserted and re.match(r'^#+\s', line):
                    new_lines.append("")
                    new_lines.append(placeholder)
                    inserted = True
            tasks.append((headline, i))
            new_slides.append("\n".join(new_lines))
        else:
            new_slides.append(slide)

    return "\n---\n".join(new_slides), tasks


def _run_image_generation(
    markdown: str | None,
    gemini_key: str,
    output_dir: str,
    parent_window,
    override_tasks: list[tuple[str, int]] | None = None,
    image_style: str = "Photorealistic",
    gemini_model: str = DEFAULT_GEMINI_MODEL,
) -> None:
    """Run bulk image generation on a background thread; posts GLib callbacks."""
    if override_tasks is None:
        GLib.idle_add(parent_window._show_toast, "Generating images in background…", 0)
    try:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            loop.run_until_complete(
                generate_all_images(
                    markdown, gemini_key, output_dir,
                    override_tasks, image_style, gemini_model,
                )
            )
        finally:
            loop.close()
            asyncio.set_event_loop(None)
    except Exception as e:
        log.warning("Image generation failed: %s", e)
        GLib.idle_add(parent_window._show_toast, f"Image generation failed: {e}", 6)
        return

    def _finish():
        parent_window._show_toast("Images ready — rebuilding slides…", 4)
        parent_window._trigger_convert()
        return GLib.SOURCE_REMOVE

    GLib.idle_add(_finish)


_MAX_DOCUMENT_BYTES = 50 * 1024 * 1024  # 50 MB


def _extract_pdf_text_from_bytes(raw: bytes) -> str:
    """Extract text from PDF bytes using Poppler then pdfminer as fallback."""
    try:
        import gi
        gi.require_version("Poppler", "0.18")
        from gi.repository import GLib as _GLib, Poppler
        gbytes = _GLib.Bytes.new(raw)
        doc = Poppler.Document.new_from_bytes(gbytes, None)
        pages = [doc.get_page(i).get_text() for i in range(doc.get_n_pages())]
        text = "\n\n".join(p for p in pages if p)
        if text.strip():
            return text
    except Exception as e:
        log.warning("Poppler PDF extraction failed: %s", e)

    import io
    try:
        from pdfminer.high_level import extract_text
        return _run_with_timeout(extract_text, io.BytesIO(raw), timeout=30)
    except Exception as exc:
        raise ImportError(
            f"Neither poppler-glib nor pdfminer.six is available for PDF extraction: {exc}"
        ) from exc


def _extract_pdf_text(path: Path) -> str:
    return _extract_pdf_text_from_bytes(path.read_bytes())


def _extract_gfile_text(gfile) -> str:
    """Extract text from a GFile — used for portal files that have no local path."""
    name = gfile.get_basename() or ""
    suffix = Path(name).suffix.lower()

    ok, raw, _ = gfile.load_contents(None)
    if not ok:
        raise ValueError("Could not read file via the document portal")
    raw = bytes(raw)

    if len(raw) > _MAX_DOCUMENT_BYTES:
        mb = len(raw) // (1024 * 1024)
        raise ValueError(f"File too large ({mb} MB > {_MAX_DOCUMENT_BYTES // (1024*1024)} MB limit)")

    if suffix in (".txt", ".md"):
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("latin-1")
    if suffix == ".pdf":
        return _extract_pdf_text_from_bytes(raw)
    if suffix == ".docx":
        import io
        try:
            import docx
        except Exception as exc:
            raise ImportError(f"python-docx not installed or failed to import: {exc}") from exc
        doc = docx.Document(io.BytesIO(raw))
        return "\n".join(p.text for p in doc.paragraphs)
    return raw.decode("utf-8", errors="replace")


def _run_with_timeout(fn, *args, timeout: float = 30.0):
    result: list = []

    def _worker():
        try:
            result.append(("ok", fn(*args)))
        except Exception as e:
            result.append(("err", e))

    t = threading.Thread(target=_worker, daemon=True)
    t.start()
    t.join(timeout=timeout)
    if t.is_alive():
        raise TimeoutError(
            f"PDF extraction timed out after {timeout}s — file may be too complex"
        )
    status, value = result[0]
    if status == "err":
        raise value
    return value


def _extract_text(path: Path) -> str:
    try:
        size = path.stat().st_size
        if size > _MAX_DOCUMENT_BYTES:
            raise ValueError(
                f"File too large ({size // (1024 * 1024)} MB > "
                f"{_MAX_DOCUMENT_BYTES // (1024 * 1024)} MB limit)"
            )
    except OSError as e:
        raise ValueError(f"Cannot read file info: {e}") from e

    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            return path.read_text(encoding="latin-1")
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise ImportError(
                "python-docx not installed — run: pip install python-docx"
            ) from exc
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    return path.read_text(encoding="utf-8")
